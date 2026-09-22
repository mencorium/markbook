# /markbook/backend/services/imports.py
"""Read class lists and mark sheets (Excel or CSV), preview the matches, then commit.
Sheets exported by this app import straight back, including per-question marks and choice sections."""
from __future__ import annotations

import csv
import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from ..phone import normalize_phone
from . import assessments as A
from . import records as R
from .analytics import Gradebook

NAME_RE = re.compile(r"^(student\s*)?(full\s*)?names?$", re.I)
MAX_RE = re.compile(r"\(/\s*(\d+(?:\.\d+)?)\)\s*$")


def norm(s) -> str:
    return " ".join(str(s or "").split()).lower()


@dataclass
class ImportRow:
    name: str
    reg: str = ""
    phone: str = ""
    mark: float | None = None
    bad: bool = False
    q: list[float | None] = field(default_factory=list)
    match_id: int | None = None
    phone_ok: bool = True


@dataclass
class ImportPreview:
    kind: str                                   # 'list' or 'marks'
    rows: list[ImportRow]
    class_name: str
    subject_name: str = ""
    assessment: str = ""
    type: str = "Test"
    date: dt.date = field(default_factory=dt.date.today)
    max_marks: float = 100
    qcols: list[dict] = field(default_factory=list)       # [{label, topic, max}]
    choices: list[dict] = field(default_factory=list)     # [{name, pick, labels}]
    file_name: str = ""


class ImportError_(ValueError):
    pass


def read_table(path: str | Path) -> list[list]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
        return [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
    with path.open(newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
        return [row for row in csv.reader(f, dialect)]


def _date(v) -> dt.date:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return dt.date.today()


def parse(rows: list[list], file_name: str, default_class: str = "") -> ImportPreview:
    rows = [[(c.strip().lstrip("\ufeff") if isinstance(c, str) else c) for c in r] for r in rows]
    h = next((i for i, r in enumerate(rows) if any(NAME_RE.match(str(c).strip()) for c in r)), None)
    if h is None:
        raise ImportError_('No "Student Name" or "Name" column header was found. Add a header row with a Name column and try again.')
    meta: dict[str, object] = {}
    for r in rows[:h]:
        c0 = r[0] if r else ""
        c1 = r[1] if len(r) > 1 else ""
        if c1 not in ("", None) and norm(c0):
            meta[norm(c0).rstrip(":")] = c1
        elif isinstance(c0, str) and ":" in c0:
            k, v = c0.split(":", 1)
            meta[norm(k)] = v.strip()
    head = [str(c).strip() for c in rows[h]]
    idx = lambda pat: next((i for i, c in enumerate(head) if re.search(pat, c, re.I)), -1)
    i_name = next(i for i, c in enumerate(head) if NAME_RE.match(c))
    i_reg, i_phone = idx(r"reg|adm|index|candidate\s*no"), idx(r"phone|mobile|simu")
    i_mark, i_grade = idx(r"^(total\s*)?(marks?|score)\b"), idx(r"^grade$")
    if i_mark < 0 and i_grade >= 0:
        raise ImportError_('This sheet has grades, not marks, so there is nothing to calculate from. Export it with "Marks" to import it.')
    qcols = []
    if i_mark >= 0:
        for i, c in enumerate(head):
            if i in (i_mark, i_name, i_reg, i_phone) or re.match(r"^total", c, re.I) or not MAX_RE.search(c):
                continue
            base = MAX_RE.sub("", c).strip()
            label, _, topic = base.partition("·")
            qcols.append({"i": i, "label": label.strip(), "topic": topic.strip(), "max": float(MAX_RE.search(c).group(1))})
    out: list[ImportRow] = []
    for r in rows[h + 1:]:
        name = str(r[i_name] if i_name < len(r) else "").strip()
        if not name:
            break
        get = lambda i: (r[i] if 0 <= i < len(r) else "")
        reg = str(get(i_reg)).strip()
        if reg.endswith(".0"):
            reg = reg[:-2]
        if reg.lower() in ("null", "n/a", "-"):
            reg = ""
        row = ImportRow(name=" ".join(name.split()), reg=reg, phone=str(get(i_phone)).strip())
        if row.phone:
            try:
                normalize_phone(row.phone)
            except ValueError:
                row.phone_ok = False
        if i_mark >= 0:
            v = get(i_mark)
            if v in ("", None) or str(v).lower().startswith("abs"):
                row.mark = None
            else:
                try:
                    row.mark = float(v)
                except (TypeError, ValueError):
                    row.bad = True
        for qc in qcols:
            v = get(qc["i"])
            try:
                row.q.append(None if v in ("", None) else float(v))
            except (TypeError, ValueError):
                row.q.append(None)
        out.append(row)
    if not out:
        raise ImportError_("The sheet has a header row but no student rows beneath it.")
    choices = []
    for part in str(meta.get("choice", "")).split(";"):
        m = re.match(r"^\s*(.*?)\s*answer any (\d+) of (.+)$", part, re.I)
        if m:
            choices.append({"name": m.group(1).strip(), "pick": int(m.group(2)), "labels": [norm(x) for x in m.group(3).split(",")]})
    asmt = str(meta.get("assessment", "")).strip()
    mx = meta.get("out of")
    head_max = MAX_RE.search(head[i_mark]).group(1) if i_mark >= 0 and MAX_RE.search(head[i_mark]) else None
    typ = next((t for t in A.TYPES if norm(t) == norm(meta.get("type"))), "Exam" if "exam" in asmt.lower() else "Test")
    return ImportPreview(
        kind="marks" if i_mark >= 0 else "list", rows=out,
        class_name=str(meta.get("class") or default_class).strip(), subject_name=str(meta.get("subject", "")).strip(),
        assessment=asmt or Path(file_name).stem, type=typ, date=_date(meta.get("date")) if meta.get("date") else dt.date.today(),
        max_marks=float(mx) if mx not in (None, "") else float(head_max) if head_max else max([100.0] + [x.mark or 0 for x in out]),
        qcols=qcols, choices=choices, file_name=Path(file_name).name)


def match_rows(p: ImportPreview, gb: Gradebook) -> ImportPreview:
    by_reg = {norm(s.reg_no): s.id for s in gb.students.values() if s.reg_no}
    cls = next((c for c in gb.classes.values() if norm(c.name) == norm(p.class_name)), None)
    by_name = {norm(s.name): s.id for s in gb.students.values() if cls and s.class_id == cls.id}
    for r in p.rows:
        r.match_id = by_reg.get(norm(r.reg)) if r.reg else None
        r.match_id = r.match_id or by_name.get(norm(r.name))
    return p


def commit(p: ImportPreview, gb: Gradebook, add_new: bool = True, subject_id: int | None = None) -> dict:
    """Creates missing students (if add_new), then for mark sheets finds or creates the assessment and writes marks."""
    if not p.class_name.strip():
        raise ImportError_("Enter a class for these students.")
    match_rows(p, gb)
    added, ids = 0, []
    for r in p.rows:
        sid = r.match_id
        if sid is None and add_new:
            sid = R.save_student(None, name=r.name, class_name=p.class_name, reg_no=r.reg or None, phone=r.phone if r.phone_ok else None).id
            added += 1
        elif sid is not None and r.phone and r.phone_ok and not gb.students[sid].phone:
            s = gb.students[sid]
            R.save_student(sid, name=s.name, class_name=gb.class_name(s.class_id), reg_no=s.reg_no, phone=r.phone)
        ids.append(sid)
    if p.kind == "list":
        return {"added": added, "marks": 0, "assessment_id": None}
    cls = R.get_or_create_class(p.class_name)
    if subject_id is None:
        subj = next((s for s in gb.subjects.values() if norm(s.name) == norm(p.subject_name) or norm(s.code) == norm(p.subject_name)), None)
        subject_id = subj.id if subj else R.save_subject(None, name=p.subject_name or "Imported subject", code=(p.subject_name or "IMP")[:4]).id
    existing = next((a for a in A.list_assessments(cls.id, subject_id) if norm(a.name) == norm(p.assessment) and a.date == p.date), None)
    if existing is None:
        a = A.save_assessment(None, subject_id=subject_id, class_id=cls.id, type=p.type, name=p.assessment, date=p.date, max_marks=p.max_marks)
        if p.qcols:
            secs = [{"key": "main", "name": "Compulsory" if p.choices else "", "pick": None}] + \
                   [{"key": f"c{j}", "name": c["name"], "pick": c["pick"]} for j, c in enumerate(p.choices)]
            qs = [{"id": None, "label": q["label"], "topic": q["topic"], "max": q["max"],
                   "section_key": next((f"c{j}" for j, c in enumerate(p.choices) if norm(q["label"]) in c["labels"]), "main")} for q in p.qcols]
            a = A.save_paper(a.id, secs, qs)
    else:
        a = existing
    n = 0
    qmap = None
    if p.qcols and a.is_paper:
        by = {norm(q.label): q for q in a.questions}
        if all(norm(c["label"]) in by for c in p.qcols):
            qmap = [by[norm(c["label"])] for c in p.qcols]
    if qmap:
        rows = {}
        for sid, r in zip(ids, p.rows):
            if sid is None:
                continue
            row = {q.id: (v if v is not None and 0 <= v <= q.max else None) for q, v in zip(qmap, r.q)}
            if any(v is not None for v in row.values()):
                rows[sid] = row
                n += 1
        A.save_question_marks(a.id, rows)
    else:
        totals = {sid: r.mark for sid, r in zip(ids, p.rows) if sid is not None and not r.bad and r.mark is not None and 0 <= r.mark <= a.max_marks}
        n = A.save_totals(a.id, totals)
    return {"added": added, "marks": n, "assessment_id": a.id}